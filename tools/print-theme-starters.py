#!/usr/bin/env python3
"""
Build the official Ekklesia PowerPoint theme starters and the export skeleton.

Run from the repository root with python-pptx installed (a development tool;
nothing here runs on the server):

    python3 tools/print-theme-starters.py

Writes resources/print/starters/*.pptx and resources/print/pptx/skeleton.pptx.
The starters are what church staff download, decorate in PowerPoint and
upload back. The contract they teach (see docs, and PptxThemeReader):

  * Slide 1 is the design. Other slides are ignored.
  * Shapes named CALENDAR_TITLE, MONTH_HEADING and CALENDAR_GRID (required)
    and SUBTITLE, LEGEND, FOOTER, PRINT_NOTE (optional) mark where Ekklesia
    writes. They are guides: Ekklesia never draws them.
  * Shapes whose name starts with GUIDE are guides too, and are ignored.
  * Everything else on slide 1 is artwork: a background, pictures, simple
    shapes and decorative text.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'resources', 'print', 'starters')
SKELETON = os.path.join(ROOT, 'resources', 'print', 'pptx')

VARIANTS = {
    'letter-portrait': (8.5, 11.0, 'Letter portrait'),
    'a4-portrait': (8.27, 11.69, 'A4 portrait'),
    'letter-landscape': (11.0, 8.5, 'Letter landscape'),
    'tabloid-portrait': (11.0, 17.0, 'Tabloid portrait (11 × 17)'),
}

INK = RGBColor(0x1F, 0x3A, 0x2E)
GUIDE = RGBColor(0x1B, 0x7F, 0xB5)
LEAF = [RGBColor(0x6F, 0x9A, 0x32), RGBColor(0xB8, 0x83, 0x22), RGBColor(0xB4, 0x4A, 0x1C), RGBColor(0x35, 0x73, 0x54)]
BAND = RGBColor(0xF3, 0xF5, 0xE3)


def guide(slide, name, x, y, w, h, label):
    """A named region: dashed outline, no fill, a label inside. Not printed by Ekklesia."""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.fill.background()
    s.line.color.rgb = GUIDE
    s.line.width = Pt(1.25)
    s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = s.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = name
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = GUIDE
    p2 = tf.add_paragraph()
    p2.text = label
    p2.alignment = PP_ALIGN.CENTER
    p2.runs[0].font.size = Pt(9)
    p2.runs[0].font.color.rgb = GUIDE
    return s


def decorate(slide, W, H):
    """Example seasonal artwork: a pale header band and leaves in two corners."""
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(1.55 if W < H else 1.35))
    band.name = 'Header band'
    band.fill.solid()
    band.fill.fore_color.rgb = BAND
    band.line.fill.background()
    leaves = [(W - 1.35, 0.15, 0.9, 0.42, 25), (W - 0.95, 0.55, 0.7, 0.32, -30), (W - 1.7, 0.62, 0.6, 0.28, 60),
              (0.2, H - 0.75, 0.9, 0.42, -20), (0.75, H - 0.55, 0.7, 0.32, 35)]
    for i, (x, y, w, h, rot) in enumerate(leaves):
        leaf = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
        leaf.name = 'Leaf %d' % (i + 1)
        leaf.rotation = rot
        leaf.fill.solid()
        leaf.fill.fore_color.rgb = LEAF[i % len(LEAF)]
        leaf.line.fill.background()


def note(slide, x, y, w, h, lines, size=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = 'GUIDE instructions'
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.runs[0].font.size = Pt(size)
        p.runs[0].font.color.rgb = INK
        if i == 0:
            p.runs[0].font.bold = True
    return tb


def starter(key, W, H, label):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    portrait = H > W

    # Slide 1: the design.
    s = prs.slides.add_slide(blank)
    decorate(s, W, H)
    m = 0.5
    # Title across the top, leaving the top-right corner for artwork; the
    # month under it on the left and the optional subtitle beside the month.
    guide(s, 'CALENDAR_TITLE', m, 0.3, W - 2 * m - 1.6, 0.72, 'The calendar title is written here')
    guide(s, 'MONTH_HEADING', m, 1.07, 3.3, 0.42, 'Month and year')
    guide(s, 'SUBTITLE', m + 3.45, 1.1, W - 2 * m - 3.45 - 1.2, 0.36, 'Optional line under the title')
    grid_top = 1.65
    grid_h = H - grid_top - (1.35 if portrait else 1.15)
    guide(s, 'CALENDAR_GRID', m, grid_top, W - 2 * m, grid_h,
          'Ekklesia draws the month here: seven days, four to six weeks. Busy weeks grow; keep this area plain.')
    # The bottom-left corner is left for artwork too.
    guide(s, 'LEGEND', m + 1.2, grid_top + grid_h + 0.1, W - 2 * m - 1.2, 0.3, 'Member-type key (optional)')
    guide(s, 'FOOTER', m + 1.2, H - 0.55, W - 2 * m - 1.2, 0.3, 'Footer (optional)')
    margin = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.25), Inches(0.25), Inches(W - 0.5), Inches(H - 0.5))
    margin.name = 'GUIDE safe margin'
    margin.fill.background()
    margin.line.color.rgb = RGBColor(0xC0, 0x3A, 0x2B)
    margin.line.width = Pt(0.75)
    margin.line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT
    # Instructions sit just off the slide: visible while editing, never printed.
    note(s, W + 0.3, 0.2, 3.6, H - 0.4, [
        'Ekklesia calendar theme — %s' % label,
        'Decorate this slide: colours, a background picture, shapes, borders and seasonal artwork.',
        'The dashed blue boxes are where Ekklesia writes. Move and resize them to suit your design, '
        'but keep their names (Home ▸ Arrange ▸ Selection Pane shows them). CALENDAR_TITLE, MONTH_HEADING '
        'and CALENDAR_GRID are required; SUBTITLE, LEGEND and FOOTER may be deleted.',
        'Keep the calendar area plain or very light: names and dates are printed over it.',
        'Keep important artwork inside the red dotted margin: most printers cannot print to the edge.',
        'Use pictures (JPEG or PNG), rectangles, rounded rectangles, ovals and text boxes. Charts, SmartArt, '
        'video, gradients and effects are not used by Ekklesia and will be left out.',
        'Do not change the slide size. Only slide 1 is used; slide 2 is an example.',
        'Save as .pptx (not .pptm) and upload it in the Print studio under PowerPoint themes.',
    ], 10)
    s.notes_slide.notes_text_frame.text = (
        'This slide is your calendar theme. The dashed blue boxes mark where Ekklesia writes the title, month, '
        'calendar grid, key and footer. Decorate around them, keep their names, and upload the .pptx to Ekklesia.')

    # Slide 2: what Ekklesia does with it.
    e = prs.slides.add_slide(blank)
    decorate(e, W, H)
    t = e.shapes.add_textbox(Inches(m), Inches(0.3), Inches(W - 2 * m - 1.6), Inches(0.72))
    t.name = 'Example title'
    t.text_frame.text = 'Birthdays'
    t.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    t.text_frame.paragraphs[0].runs[0].font.bold = True
    t.text_frame.paragraphs[0].runs[0].font.color.rgb = INK
    mo = e.shapes.add_textbox(Inches(m), Inches(1.05), Inches(3.3), Inches(0.45))
    mo.name = 'Example month'
    mo.text_frame.text = 'September 2026'
    mo.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    rows, cols = 5, 7
    gw, gh = W - 2 * m, grid_h
    for r in range(rows):
        for c in range(cols):
            cell = e.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(m + c * gw / cols), Inches(grid_top + r * gh / rows),
                                      Inches(gw / cols), Inches(gh / rows))
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            cell.line.color.rgb = RGBColor(0xD3, 0xDD, 0xD7)
            cell.line.width = Pt(0.5)
            tf = cell.text_frame
            tf.text = str(r * cols + c + 1) if r * cols + c < 30 else ''
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT
            if tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].font.size = Pt(9)
                tf.paragraphs[0].runs[0].font.color.rgb = INK
    note(e, m, H - 0.5, W - 2 * m, 0.4, ['Example only — slide 2 is not imported. Ekklesia fills slide 1 with your real calendar.'], 9)

    prs.core_properties.title = 'Ekklesia calendar theme'
    prs.core_properties.subject = 'Ekklesia calendar theme starter (%s)' % label
    prs.core_properties.keywords = 'ekklesia-theme; contract v1'
    prs.core_properties.author = 'Ekklesia'
    prs.save(os.path.join(OUT, 'ekklesia-calendar-theme-%s.pptx' % key))


def skeleton():
    """An empty, PowerPoint-authored package: master, layouts and theme. The
    editable export (PptxCalendarWriter) adds slides to it."""
    os.makedirs(SKELETON, exist_ok=True)
    prs = Presentation()
    prs.core_properties.title = 'Ekklesia calendar'
    prs.core_properties.author = 'Ekklesia'
    prs.save(os.path.join(SKELETON, 'skeleton.pptx'))


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    for key, (w, h, label) in VARIANTS.items():
        starter(key, w, h, label)
    skeleton()
    print('starters and skeleton written')
